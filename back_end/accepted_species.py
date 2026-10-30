"""Keeps the spreadsheet of accepted species: one row per sample."""
from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADER = ["Sample", "Genus", "Species"]


def split_species(full_name: str) -> tuple[str, str]:
    """'Penicillium glabrum' -> ('Penicillium', 'glabrum')."""
    parts = full_name.strip().split(None, 1)
    genus = parts[0] if parts else ""
    epithet = parts[1] if len(parts) > 1 else "sp."
    return genus, epithet


def add_accepted(xlsx_path: Path, sample: str, full_species_name: str) -> bool:
    """Record an accepted hit. Returns False if it was already recorded.

    New sample  -> new row: sample, genus, species.
    Known sample -> genus is kept, the species goes in the next free column.
    """
    genus, epithet = split_species(full_species_name)

    if xlsx_path.exists():
        wb = load_workbook(xlsx_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Accepted"
        ws.append(HEADER)

    # Find this sample's row (row 1 is the header)
    row_idx = next(
        (r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == sample),
        None,
    )

    if row_idx is None:
        ws.append([sample, genus, epithet])
        added = True
    else:
        existing = [
            str(ws.cell(row_idx, c).value).strip().lower()
            for c in range(3, ws.max_column + 1)
            if ws.cell(row_idx, c).value not in (None, "")
        ]
        if epithet.lower() in existing:
            added = False
        else:
            col = 3
            while ws.cell(row_idx, col).value not in (None, ""):
                col += 1
            ws.cell(row_idx, col, epithet)
            added = True

    if added:
        wb.save(xlsx_path)      # raises PermissionError if the file is open in Excel
    return added